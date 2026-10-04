"""Render the editable JSON manuscript into Korean PDF, DOCX, Typst and HTML."""
from pathlib import Path
import json,html,sys
import typst
from docx import Document
from docx.shared import Pt,Inches,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];P=ROOT/'paper'
doc=json.loads((P/'manuscript.json').read_text())
q=lambda s:json.dumps(str(s),ensure_ascii=False)
txt=lambda s:'#text('+q(s)+')'
t=['#set document(title: '+q(doc['title'].replace('\n',' '))+', author: "Research draft")',
   '#set page(paper: "a4", margin: (x: 19mm, top: 18mm, bottom: 19mm), header: align(right, text(size: 8pt, fill: rgb("666666"))[Hydraulic-pump anomaly detection · Korean research draft]), footer: align(center, context counter(page).display("1")))',
   '#set text(font: "Noto Sans KR", size: 9.6pt, lang: "ko")',
   '#set par(justify: true, leading: 0.65em, spacing: 0.85em)',
   '#set heading(outlined: true)',
   '#show heading.where(level: 1): set text(size: 14pt, weight: "bold", fill: rgb("123B55"))',
   '#show heading.where(level: 2): set text(size: 11pt, weight: "bold")',
   '#set table(stroke: 0.35pt + rgb("CBD5E1"), inset: 4pt)',
   '#align(center)[#text(size: 19pt, weight: "bold")['+txt(doc['title'])+']\n\n'+txt(doc['english_title'])+'\n\n#text(size: 9pt)['+txt(doc['author'])+'\n\n'+txt(doc['version'])+']]\n#v(6mm)']
w=Document(); sec=w.sections[0];sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
for style in ['Normal','Title','Heading 1','Heading 2','Caption']:
    s=w.styles[style];s.font.name='Noto Sans KR';s._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans KR')
w.styles['Normal'].font.size=Pt(10);w.styles['Normal'].paragraph_format.line_spacing=1.2
w.add_heading(doc['title'],0);w.add_paragraph(doc['english_title']);w.add_paragraph(doc['author']);w.add_paragraph(doc['version'])
hf=['<!doctype html><html lang="ko"><meta charset="utf-8"><title>'+html.escape(doc['title'])+'</title><style>body{font-family:system-ui,sans-serif;max-width:1000px;margin:40px auto;padding:20px;line-height:1.8;color:#182333}h1,h2{color:#123b55}table{border-collapse:collapse;font-size:13px;width:100%;margin:15px 0}td,th{border:1px solid #cbd5e1;padding:6px}th{background:#eef3f7}img{max-width:100%}.caption{font-size:13px;color:#445}</style><h1>'+html.escape(doc['title']).replace('\n','<br>')+'</h1><p>'+html.escape(doc['english_title'])+'</p><p>'+html.escape(doc['author'])+'</p>']
plain=[doc['title'],doc['english_title'],doc['author']]
for b in doc['blocks']:
    kind=b['type']
    if kind=='heading':
        t.append('#heading(level: '+str(b['level'])+')['+txt(b['text'])+']');w.add_heading(b['text'],b['level']);hf.append(f'<h{b["level"]+1}>{html.escape(b["text"])}</h{b["level"]+1}>');plain.append(b['text'])
    elif kind in ['paragraph','equation']:
        if kind=='equation':t.append('#block(fill: rgb("F2F5F8"), inset: 8pt, width: 100%)['+txt(b['text'])+']')
        else:t.append(txt(b['text'])+'\n')
        w.add_paragraph(b['text']);hf.append('<p>'+html.escape(b['text'])+'</p>');plain.append(b['text'])
    elif kind=='table':
        n=len(b['headers']);cells=['['+txt(v)+']' for row in b['rows'] for v in row]
        headers=','.join('['+txt(v)+']' for v in b['headers'])
        # Rows can flow across pages. Header is repeated by Typst and Word.
        t.append('#text(size: 8.4pt, weight: "bold")['+txt(b['caption'])+']\n#block[')
        t.append('#set text(size: '+('7.0pt' if n>=9 else '7.6pt')+')\n#table(columns: '+str(n)+', table.header('+headers+'),'+','.join(cells)+')\n]')
        w.add_paragraph(b['caption'],'Caption');table=w.add_table(rows=1,cols=n);table.style='Table Grid'
        for i,v in enumerate(b['headers']):table.rows[0].cells[i].text=v
        trpr=table.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for row in b['rows']:
            cellsw=table.add_row().cells
            for i,v in enumerate(row):cellsw[i].text=v
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:run.font.size=Pt(7.5)
        hf.append('<p class="caption">'+html.escape(b['caption'])+'</p><table><thead><tr>'+''.join('<th>'+html.escape(v)+'</th>' for v in b['headers'])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+html.escape(v)+'</td>' for v in r)+'</tr>' for r in b['rows'])+'</tbody></table>')
        plain+=[b['caption'],' | '.join(b['headers'])]+[' | '.join(row) for row in b['rows']]
    elif kind=='image':
        t.append('#figure(image('+q(b['path'])+', width: 100%), caption: none)\n#text(size: 8.4pt)['+txt(b['caption'])+']')
        w.add_picture(str(P/b['path']),width=Cm(16.8));w.add_paragraph(b['caption'],'Caption')
        hf.append('<figure><img src="'+html.escape(b['path'])+'"><figcaption>'+html.escape(b['caption'])+'</figcaption></figure>');plain.append(b['caption'])
    elif kind=='pagebreak':t.append('#pagebreak()');w.add_page_break()
source=P/'manuscript.typ';source.write_text('\n\n'.join(t))
typst.compile(str(source),output=str(P/'논문_한국어.pdf'),font_paths=[str(P/'fonts')],root=str(ROOT))
w.save(P/'논문_한국어.docx');(P/'논문_한국어.html').write_text(''.join(hf)+'</html>');(P/'논문_한국어.txt').write_text('\n\n'.join(plain))
from pypdf import PdfReader
r=PdfReader(P/'논문_한국어.pdf');extracted='\n'.join(p.extract_text() for p in r.pages)
for needle in ['303','0.8487','0.9182','부록','참고문헌','관측']:
    assert needle in extracted,needle
info=dict(pages=len(r.pages),pdf_bytes=(P/'논문_한국어.pdf').stat().st_size,typst_version=typst.__version__,python=sys.version,checked_text=['303','0.8487','0.9182','부록','참고문헌','관측'])
(P/'build_verification.json').write_text(json.dumps(info,ensure_ascii=False,indent=2));print(json.dumps(info,ensure_ascii=False))
